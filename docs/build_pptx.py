# -*- coding: utf-8 -*-
"""Sinh file PowerPoint giới thiệu sản phẩm (tiếng Nhật) từ nội dung README."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- Bảng màu thương hiệu ----
ORANGE = RGBColor(0xF3, 0x80, 0x20)   # Cloudflare orange
DARK = RGBColor(0x18, 0x18, 0x1B)
GRAY = RGBColor(0x52, 0x52, 0x5B)
LIGHT = RGBColor(0xFA, 0xFA, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
ROW_ALT = RGBColor(0xF4, 0xF4, 0xF5)

FONT = "Yu Gothic UI"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    rect = s.shapes.add_shape(1, 0, 0, SW, SH)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    s.shapes._spTree.remove(rect._element)
    s.shapes._spTree.insert(2, rect._element)
    return s


def txt(slide, left, top, width, height, text, size=18, color=DARK, bold=False,
        align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = ln
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return tb


def title_bar(slide, text, num=None):
    txt(slide, Inches(0.6), Inches(0.45), Inches(11.5), Inches(0.9), text,
        size=32, color=DARK, bold=True)
    line = slide.shapes.add_shape(1, Inches(0.62), Inches(1.32), Inches(3.2), Pt(4))
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.fill.background()
    line.shadow.inherit = False
    if num:
        txt(slide, Inches(12.3), Inches(0.5), Inches(0.8), Inches(0.5), num,
            size=14, color=GRAY, align=PP_ALIGN.RIGHT)


def bullets(slide, left, top, width, items, size=18, gap=1.45, color=DARK):
    tb = slide.shapes.add_textbox(left, top, width, Inches(5))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (mark, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = gap
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = mark + "  "
        r1.font.size = Pt(size)
        r1.font.name = FONT
        r1.font.color.rgb = ORANGE
        r1.font.bold = True
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(size)
        r2.font.name = FONT
        r2.font.color.rgb = color
    return tb


def make_table(slide, left, top, width, rows, col_widths=None, header_fill=DARK,
               font_size=14, header_size=14):
    nrows = len(rows)
    ncols = len(rows[0])
    height = Inches(0.5 * nrows)
    gtbl = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    tbl = gtbl.table
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = w
    for r in range(nrows):
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = rows[r][c]
            run.font.name = FONT
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
                run.font.color.rgb = WHITE
                run.font.bold = True
                run.font.size = Pt(header_size)
                if c > 0:
                    p.alignment = PP_ALIGN.CENTER
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r % 2 else ROW_ALT
                run.font.color.rgb = DARK
                run.font.size = Pt(font_size)
                if ncols > 2 and c > 0:
                    p.alignment = PP_ALIGN.CENTER
    return gtbl


def code_box(slide, left, top, width, height, text, size=12):
    box = slide.shapes.add_shape(1, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x1F, 0x1F, 0x23)
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.18)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = ln
        run.font.size = Pt(size)
        run.font.name = MONO
        run.font.color.rgb = RGBColor(0xE4, 0xE4, 0xE7)
    return box


def node(slide, left, top, w, h, title, sub=None, fill=WHITE, border=ORANGE,
         title_color=DARK, title_size=15, sub_size=10.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(title_size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = title_color
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(sub_size)
        r2.font.name = FONT
        r2.font.color.rgb = GRAY
    return shp


def arrow(slide, shape_type, left, top, w, h, color=RGBColor(0x9C, 0xA3, 0xAF)):
    a = slide.shapes.add_shape(shape_type, left, top, w, h)
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


# ============================================================
# Slide 1 — Title
# ============================================================
s = add_slide(DARK)
band = s.shapes.add_shape(1, 0, Inches(3.05), SW, Inches(0.12))
band.fill.solid(); band.fill.fore_color.rgb = ORANGE; band.line.fill.background(); band.shadow.inherit = False
txt(s, Inches(1), Inches(1.6), Inches(11.3), Inches(1.2), "📋 Task Management Dashboard",
    size=46, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(3.35), Inches(11.3), Inches(0.8),
    "プロジェクトとタスクをチームで管理するフルスタック Web アプリ",
    size=22, color=RGBColor(0xD4, 0xD4, 0xD8), align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(4.6), Inches(11.3), Inches(0.6),
    "Next.js 15 · Cloudflare Workers · D1 · R2 · Firebase Auth · Resend",
    size=16, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(6.4), Inches(11.3), Inches(0.5),
    "製品紹介プレゼンテーション — 2026年", size=13, color=GRAY, align=PP_ALIGN.CENTER)

# ============================================================
# Slide 2 — 製品概要
# ============================================================
s = add_slide()
title_bar(s, "🎯 製品概要", "02")
txt(s, Inches(0.62), Inches(1.7), Inches(11.5), Inches(0.6),
    "認証済みユーザーが自分の仕事を整理・追跡できる生産性ツールです。",
    size=19, color=GRAY)
bullets(s, Inches(0.7), Inches(2.6), Inches(11.8), [
    ("●", "仕事を プロジェクト（Projects）単位でまとめる"),
    ("●", "プロジェクトを追跡可能な タスク（Tasks）に分解"),
    ("●", "チーム（Teams）で共同作業し、メンバーをメールで招待"),
    ("●", "ステータス・優先度・期限・画像添付に対応"),
    ("●", "エッジで高速・低コストに動作（常時稼働サーバー不要）"),
], size=19)

# ============================================================
# Slide 3 — 主な機能
# ============================================================
s = add_slide()
title_bar(s, "✨ 主な機能", "03")
make_table(s, Inches(0.62), Inches(1.7), Inches(12.1), [
    ["カテゴリ", "内容"],
    ["🔐 認証", "登録・ログイン・メール認証・パスワードリセット（Firebase）"],
    ["📁 プロジェクト", "作成 / 編集 / 削除、サムネイル画像、進捗統計"],
    ["✅ タスク", "ステータス・優先度・担当者・タグ・期限・画像・コメント"],
    ["👥 チーム", "チーム作成、メンバー管理、メール招待、プロジェクト割当"],
    ["🖼️ 画像", "プロジェクト / タスク / アバターを R2 に保存"],
    ["🌐 多言語", "英語（既定）・日本語・ベトナム語の3言語対応"],
], col_widths=[Inches(3.2), Inches(8.9)], font_size=15)

# ============================================================
# Slide 4 — アーキテクチャ
# ============================================================
s = add_slide()
title_bar(s, "🏗️ アーキテクチャ", "04")

# 上段：フロント ＋ Worker
node(s, Inches(0.7), Inches(1.95), Inches(4.4), Inches(1.5),
     "Next.js 15（フロントエンド）", "app/ · components/ · Firebase Web SDK",
     border=ACCENT, title_size=15)
node(s, Inches(8.25), Inches(1.95), Inches(4.4), Inches(1.5),
     "Cloudflare Worker (Hono)", "Firebase ID JWT を検証 · REST API",
     border=ORANGE, title_size=15)
# 横矢印（HTTPS / CORS）
txt(s, Inches(5.2), Inches(2.05), Inches(2.85), Inches(0.4),
    "HTTPS / CORS", size=12, color=GRAY, align=PP_ALIGN.CENTER, bold=True)
arrow(s, MSO_SHAPE.RIGHT_ARROW, Inches(5.25), Inches(2.5), Inches(2.75), Inches(0.45),
      color=ORANGE)
txt(s, Inches(5.2), Inches(3.0), Inches(2.85), Inches(0.4),
    "ID トークン", size=11, color=GRAY, align=PP_ALIGN.CENTER)

# 縦矢印
arrow(s, MSO_SHAPE.DOWN_ARROW, Inches(2.67), Inches(3.55), Inches(0.45), Inches(1.2))
arrow(s, MSO_SHAPE.DOWN_ARROW, Inches(8.75), Inches(3.55), Inches(0.4), Inches(1.2))
arrow(s, MSO_SHAPE.DOWN_ARROW, Inches(10.25), Inches(3.55), Inches(0.4), Inches(1.2))
arrow(s, MSO_SHAPE.DOWN_ARROW, Inches(11.78), Inches(3.55), Inches(0.4), Inches(1.2))

# 下段：サービス
node(s, Inches(0.7), Inches(4.85), Inches(4.4), Inches(1.35),
     "Firebase Auth", "認証（マネージド）", border=RGBColor(0xFF, 0xCA, 0x28),
     title_size=15)
node(s, Inches(8.25), Inches(4.85), Inches(1.3), Inches(1.35),
     "D1", "SQLite", border=ORANGE, title_size=14)
node(s, Inches(9.75), Inches(4.85), Inches(1.3), Inches(1.35),
     "R2", "画像", border=ORANGE, title_size=14)
node(s, Inches(11.25), Inches(4.85), Inches(1.4), Inches(1.35),
     "Resend", "メール", border=ORANGE, title_size=14)

txt(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.6),
    "フロントは ID トークン付きで Worker を呼び出し、Worker が D1 / R2 / Resend を制御します。",
    size=14, color=GRAY)

# ============================================================
# Slide 5 — 技術スタック
# ============================================================
s = add_slide()
title_bar(s, "🧰 技術スタック", "05")
make_table(s, Inches(0.62), Inches(1.7), Inches(12.1), [
    ["レイヤー", "技術"],
    ["フロントエンド", "Next.js 15 (App Router) + React 19"],
    ["UI", "Tailwind CSS + Radix UI + lucide-react"],
    ["データ取得", "TanStack Query"],
    ["フォーム / 検証", "React Hook Form + Zod"],
    ["認証（クライアント）", "Firebase Web SDK"],
    ["API ランタイム", "Cloudflare Workers + Hono"],
    ["データベース", "Cloudflare D1 (SQLite)"],
    ["ファイル保存", "Cloudflare R2"],
    ["メール送信", "Resend (HTTP API)"],
], col_widths=[Inches(3.6), Inches(8.5)], font_size=14)

# ============================================================
# Slide 6 — チーム機能
# ============================================================
s = add_slide()
title_bar(s, "👥 注目機能：チーム（Teams）", "06")
txt(s, Inches(0.62), Inches(1.65), Inches(11.8), Inches(0.6),
    "仕事をチーム単位で共同管理できる、今回の目玉機能です。",
    size=18, color=GRAY)
bullets(s, Inches(0.7), Inches(2.5), Inches(12), [
    ("➕", "サイドバーからワンクリックでチーム作成（画像・公開範囲）"),
    ("📜", "チーム一覧（無限スクロールで大量データも快適）"),
    ("🔍", "チーム詳細：メンバー・役割・プロジェクトを一覧表示"),
    ("📧", "メールで招待（Resend）→ リンクから参加"),
    ("📁", "既存プロジェクトをチームに割り当て"),
    ("🗑️", "キーワード入力による安全な削除"),
], size=18, gap=1.4)

# ============================================================
# Slide 7 — メール招待フロー
# ============================================================
s = add_slide()
title_bar(s, "📧 メール招待フロー", "07")

steps = [
    ("① オーナー / 管理者がメールを入力", None, WHITE, ORANGE),
    ("② 招待を DB に保存 ＋ Resend でメール送信", None, WHITE, ORANGE),
    ("③ 受信者がメール内リンクをクリック", "/invite?token=...", WHITE, ORANGE),
    ("④ ログイン（招待メールとの一致を確認）", None, WHITE, ORANGE),
    ("⑤ 「承諾」→ チームメンバーに追加 ✅", None, RGBColor(0xFF, 0xF2, 0xE6), ORANGE),
]
left = Inches(0.7)
bw = Inches(7.5)
bh = Inches(0.78)
cx = left + Emu(int(bw) // 2)
top = Inches(1.7)
step_gap = Inches(1.0)
for i, (title, sub, fill, border) in enumerate(steps):
    t = Emu(int(top) + int(step_gap) * i)
    node(s, left, t, bw, bh, title, sub, fill=fill, border=border, title_size=14, sub_size=11)
    if i < len(steps) - 1:
        arrow(s, MSO_SHAPE.DOWN_ARROW,
              Emu(int(cx) - int(Inches(0.18))), Emu(int(t) + int(bh) + int(Inches(0.04))),
              Inches(0.36), Emu(int(step_gap) - int(bh) - int(Inches(0.08))))

# 注意ボックス（右側）
note = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(8.7), Inches(3.0), Inches(4.0), Inches(2.0))
note.fill.solid()
note.fill.fore_color.rgb = RGBColor(0xEF, 0xF6, 0xFF)
note.line.color.rgb = ACCENT
note.line.width = Pt(1.5)
note.shadow.inherit = False
tf = note.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.25)
tf.margin_right = Inches(0.25)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "🔒 セキュリティ"
r.font.size = Pt(15)
r.font.bold = True
r.font.name = FONT
r.font.color.rgb = ACCENT
p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.line_spacing = 1.25
p2.space_before = Pt(8)
r2 = p2.add_run()
r2.text = "招待メールとログイン中のアカウントが一致しないと承諾不可（なりすまし防止）"
r2.font.size = Pt(13)
r2.font.name = FONT
r2.font.color.rgb = DARK

# ============================================================
# Slide 8 — 役割と権限
# ============================================================
s = add_slide()
title_bar(s, "🔐 役割と権限", "08")
make_table(s, Inches(0.62), Inches(1.7), Inches(12.1), [
    ["操作", "owner", "admin", "member"],
    ["閲覧（チーム / メンバー / プロジェクト）", "✅", "✅", "✅"],
    ["招待 / 招待取消", "✅", "✅", "✕"],
    ["役割の変更", "✅", "✕", "✕"],
    ["メンバー削除", "✅", "✅", "✕"],
    ["プロジェクト割当 / 解除", "✅", "✅", "✕"],
    ["チーム削除", "✅", "✕", "✕"],
], col_widths=[Inches(6.1), Inches(2.0), Inches(2.0), Inches(2.0)], font_size=15)
txt(s, Inches(0.62), Inches(6.2), Inches(12), Inches(0.6),
    "※ 権限はフロント（UI 非表示）とバックエンド（API 検証）の両方でチェック",
    size=14, color=GRAY)

# ============================================================
# Slide 9 — パフォーマンス
# ============================================================
s = add_slide()
title_bar(s, "⚡ パフォーマンスへの工夫", "09")
txt(s, Inches(0.62), Inches(1.65), Inches(11.8), Inches(0.6),
    "Keyset（カーソル）ページネーションを採用", size=20, color=DARK, bold=True)
bullets(s, Inches(0.7), Inches(2.5), Inches(12), [
    ("▸", "OFFSET を使わず、(updated_at, id) をカーソルに"),
    ("▸", "インデックス idx_teams_updated を常に活用"),
    ("▸", "チームが 10件でも 100万件でも速度が一定"),
    ("▸", "一度に20件だけ取得し、スクロール時に追加読み込み"),
], size=17, gap=1.35)
code_box(s, Inches(0.62), Inches(5.3), Inches(12.1), Inches(1.4),
    "WHERE (updated_at < ? OR (updated_at = ? AND id < ?))\n"
    "ORDER BY updated_at DESC, id DESC  LIMIT 21", size=14)

# ============================================================
# Slide 10 — 多言語対応
# ============================================================
s = add_slide()
title_bar(s, "🌐 多言語対応（i18n）", "10")
txt(s, Inches(0.62), Inches(1.65), Inches(11.8), Inches(0.6),
    "すべての文字列を辞書で管理し、ハードコードなし", size=19, color=GRAY)
make_table(s, Inches(0.62), Inches(2.5), Inches(9.5), [
    ["言語", "コード", "状態"],
    ["English", "en", "既定"],
    ["日本語", "ja", "✅ 完全対応"],
    ["Tiếng Việt", "vi", "✅ 完全対応"],
], col_widths=[Inches(4.0), Inches(2.5), Inches(3.0)], font_size=15)
bullets(s, Inches(0.7), Inches(5.2), Inches(12), [
    ("▸", "型安全な翻訳キー（TypeScript が漏れを検出）"),
    ("▸", "削除確認キーワードも言語別：XOA / DELETE / 削除"),
], size=16, gap=1.3)

# ============================================================
# Slide 11 — デプロイ構成
# ============================================================
s = add_slide()
title_bar(s, "🚀 デプロイ構成", "11")
make_table(s, Inches(0.62), Inches(1.7), Inches(12.1), [
    ["部分", "ホスティング"],
    ["フロントエンド", "Vercel / Netlify / Cloudflare Pages"],
    ["API (Worker)", "Cloudflare Workers（エッジ）"],
    ["データベース", "Cloudflare D1"],
    ["ファイル", "Cloudflare R2"],
    ["認証", "Firebase（マネージド）"],
    ["メール", "Resend"],
], col_widths=[Inches(3.6), Inches(8.5)], font_size=15)
band = s.shapes.add_shape(1, Inches(0.62), Inches(6.05), Inches(12.1), Inches(0.85))
band.fill.solid(); band.fill.fore_color.rgb = ORANGE; band.line.fill.background(); band.shadow.inherit = False
tf = band.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "✅ 常時稼働サーバー不要・エグレス料金なし・グローバルに高速"
r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT

# ============================================================
# Slide 12 — まとめ
# ============================================================
s = add_slide(DARK)
band = s.shapes.add_shape(1, 0, Inches(3.05), SW, Inches(0.12))
band.fill.solid(); band.fill.fore_color.rgb = ORANGE; band.line.fill.background(); band.shadow.inherit = False
txt(s, Inches(1), Inches(2.0), Inches(11.3), Inches(1),
    "🙏 ありがとうございました", size=42, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(3.45), Inches(11.3), Inches(0.7),
    "Task Management Dashboard", size=22, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(4.5), Inches(11.3), Inches(0.7),
    "プロジェクト × タスク × チームを、ひとつのダッシュボードで。",
    size=18, color=RGBColor(0xD4, 0xD4, 0xD8), align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(6.4), Inches(11.3), Inches(0.5),
    "詳細ドキュメント: README.md / docs/teams-feature.md",
    size=13, color=GRAY, align=PP_ALIGN.CENTER)

out = "製品紹介_TaskManagement_JP.pptx"
prs.save(out)
print("Saved OK -", len(prs.slides._sldIdLst), "slides")
